"""Makes the Word draft templates in docs/templates/ (PRD Section 4, Phase 6), one per category,
in the site's house style: Georgia headings, Calibri body, navy and gold, branded header and footer,
borderless modern tables. Run: python scripts/make-templates.py
"""
import os
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

NAVY = RGBColor(0x0C, 0x44, 0x7C)
BLUE = RGBColor(0x18, 0x5F, 0xA5)
GOLD = RGBColor(0xC9, 0xA8, 0x4C)
GREY = RGBColor(0x4F, 0x61, 0x76)
HINT = RGBColor(0x6B, 0x7A, 0x8C)

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'docs', 'templates')

PRIVACY = ("Never put private information in a post, even in a draft: no student names, phone numbers, "
           "fees paid by a family, or internal notes. The website's files are public.")

COMMON_FIELDS = [
    ("Title", "Plain and clear. For example: Becoming a doctor: MBBS in India and abroad"),
    ("Summary (max 160 characters)", "One or two sentences for the card and search results."),
    ("Also show in (optional, up to 2)", "Careers / Internships & Opportunities / Courses / Countries / Universities"),
    ("Grade band (pick one or more)", "Grades 6 to 8 / Grades 9 and 10 / Grades 11 and 12"),
    ("Stream (optional)", "Science / Commerce / Humanities / Any stream"),
    ("Tags: Subject area (at least 1)", "For example: Engineering, Medicine, Data and AI, Studying abroad"),
    ("Tags: Entrance exam (optional)", "For example: JEE, NEET, SAT, IELTS and TOEFL"),
    ("Tags: Skill (optional)", "For example: Coding, Research, Writing, Study habits"),
    ("Tags: Study level (optional)", "School / Undergraduate / Postgraduate / Diploma and vocational"),
    ("Tags: Audience (optional)", "Students / Parents / Middle school. 3 to 8 tags in total across all groups."),
    ("Countries (optional)", "For example: United Kingdom, Canada"),
    ("For Parents box (2 to 4 sentences)", "What this is, rough cost, the key date, and one question for the family to ask a career counsellor."),
    ("Official links", "Label and web address of each official page the reader should check."),
    ("Key dates (optional)", "Each important date in the post: the date, what happens, a note such as the time zone, and the official page."),
    ("Cover image (optional)", "Only free stock (Unsplash, Pexels) or your own photos. Add a one line description and the credit."),
    ("Written by", "Your name"),
    ("Checked by", "The counsellor who checked every fact, figure and link"),
    ("Last checked", "Date, for example 11 October 2026"),
]

OPPORTUNITY_FIELDS = [
    ("Type", "Internship / Summer programme / Competition / Webinar / Scholarship / Other"),
    ("Deadline", "Last date to apply. It is treated as 11:59 PM India time. Leave empty if rolling."),
    ("Deadline note (optional)", "For example: Closes 11:59 PM US Eastern time, which is the next morning in India"),
    ("Rolling (no fixed deadline)?", "Yes / No"),
    ("Event dates (optional)", "For example: 15 to 30 June 2027"),
    ("Mode", "Online / In person / Hybrid"),
    ("Cost", "Free / Paid"),
    ("Apply link", "The official page where a student applies"),
]

BODIES = {
    'careers': ("Careers", [
        ("What is this career, in simple words", "Two or three short paragraphs a Grade 6 student can follow."),
        ("A day in the life", "What a typical day looks like. Bullets are fine."),
        ("Is it a good fit for you? (interests and strengths)", "Who enjoys this work, and who may not."),
        ("What to study in school (subjects and stream)", "Subjects for Grades 9 to 12, for both CBSE and Cambridge students."),
        ("Courses that lead here (India and abroad)", "Degree names, length, and where they are offered."),
        ("Entrance tests", "Name each test and link its official page."),
        ("Skills to start building now", "Things a student can do this year."),
        ("Where people work and how the career grows", "Employers, sectors, and the usual path upward."),
        ("Approximate starting pay in India (marked \"approximate\")", "Give a range and say it is approximate."),
        ("Related careers", "Two to four nearby careers, with a line on each."),
    ]),
    'courses': ("Courses", [
        ("What the course is about", ""),
        ("Who it suits", ""),
        ("Subjects needed in Grades 11 and 12", ""),
        ("Entrance tests and admission routes", "Name each test and link its official page."),
        ("Good places to study it (India and abroad)", ""),
        ("Duration and approximate cost", "Mark every figure approximate and give the year."),
        ("Careers it leads to", ""),
    ]),
    'opportunities': ("Internships & Opportunities", [
        ("What it is", ""),
        ("Who can apply (grade, age, subjects)", ""),
        ("What you will do or gain", ""),
        ("How to apply, step by step", "Numbered steps with dates."),
        ("Cost and any scholarship", "Mark figures approximate; link the official fee page."),
        ("Tips from the author", ""),
    ]),
    'countries': ("Countries", [
        ("Why students choose this country", ""),
        ("Popular courses", ""),
        ("Approximate cost of study and living", "A small table: tuition range, living costs, with the year and currency."),
        ("Tests needed (English and others)", ""),
        ("Application timeline", "A small table: when, what."),
        ("Visa basics (with link to the official site)", "Link only the government site."),
        ("Scholarships", ""),
        ("Working while studying and after", ""),
        ("Things to think about (climate, distance, safety)", ""),
    ]),
    'universities': ("Universities", [
        ("At a glance", "University, location, founded, popular subjects, boards accepted, medium of instruction, latest rankings with the year."),
        ("Take the virtual tour", "Link to an official or independent campus tour video."),
        ("Admission requirements for CBSE and Cambridge A Level", "A table with three columns: Requirement, CBSE Grade 12, Cambridge A Level."),
        ("Common to both pathways", ""),
        ("Voices from the university", "Only quotes you can check on an official page, with the source."),
        ("Why should your child consider this university?", "Three or four numbered reasons."),
        ("What parents should know", "Bullets. Start the most important ones with CRITICAL."),
        ("Talk to me", "Invite families to write to pcmadankumar@gmail.com or use the contact page."),
        ("Sources", "Numbered list: what the page says, and its web address."),
    ]),
}


def set_cell_borders_none(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'nil')
        borders.append(el)
    tcPr.append(borders)


def set_bottom_border(cell, color='D6E2F0', size='4'):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'right'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'nil')
        borders.append(el)
    b = OxmlElement('w:bottom')
    b.set(qn('w:val'), 'single')
    b.set(qn('w:sz'), size)
    b.set(qn('w:color'), color)
    borders.append(b)
    tcPr.append(borders)


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def run(par, text, bold=False, italic=False, size=11, color=None, font='Calibri'):
    r = par.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    r.font.name = font
    r._element.rPr.rFonts.set(qn('w:eastAsia'), font)
    if color is not None:
        r.font.color.rgb = color
    return r


def heading(doc, text, level=1):
    p = doc.add_paragraph()
    size = {1: 20, 2: 14, 3: 12}[level]
    run(p, text, bold=True, size=size, color=NAVY, font='Georgia')
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(4)
    return p


def hint(doc, text):
    p = doc.add_paragraph()
    run(p, text, italic=True, size=10, color=HINT)
    p.paragraph_format.space_after = Pt(2)
    return p


def body_placeholder(doc):
    p = doc.add_paragraph()
    run(p, "Write here.", size=11, color=HINT)
    p.paragraph_format.space_after = Pt(10)


def fields_table(doc, rows):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    hdr = table.rows[0].cells
    for cell, text in zip(hdr, ("Field", "What to write")):
        cell.text = ""
        run(cell.paragraphs[0], text, bold=True, size=11, color=NAVY, font='Georgia')
        set_bottom_border(cell, '0C447C', '12')
    for label, note in rows:
        cells = table.add_row().cells
        cells[0].text = ""
        run(cells[0].paragraphs[0], label, bold=True, size=10.5)
        cells[1].text = ""
        run(cells[1].paragraphs[0], note, size=10, color=GREY)
        for c in cells:
            set_bottom_border(c)
    # column widths
    for row in table.rows:
        row.cells[0].width = Cm(5.5)
        row.cells[1].width = Cm(11)
    doc.add_paragraph()


def make(key, label, sections):
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(2)
        s.bottom_margin = Cm(2)
        s.left_margin = Cm(2.2)
        s.right_margin = Cm(2.2)
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')

    # header: navy wordmark, gold rule
    section = doc.sections[0]
    hp = section.header.paragraphs[0]
    run(hp, "Future Pathways", bold=True, size=14, color=NAVY, font='Georgia')
    run(hp, f"   ·   Post draft template: {label}", size=10, color=GREY)
    rule = section.header.add_paragraph()
    pPr = rule._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '12')
    bottom.set(qn('w:color'), 'C9A84C')
    pbdr.append(bottom)
    pPr.append(pbdr)
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(fp, "Future Pathways, by Madan Kumar PC  ·  pcmadankumar@gmail.com  ·  "
            "When the draft is ready, send it to Madan.", size=9, color=GREY)

    # title block
    p = doc.add_paragraph()
    run(p, f"{label} post: draft", bold=True, size=22, color=NAVY, font='Georgia')
    p = doc.add_paragraph()
    run(p, "Fill in the details table, then write under each heading. Skip a heading if it does not apply. "
           "Write for a Grade 6 student and their parent: short sentences, no jargon. Mark every fee, salary, "
           "cut-off and deadline \"approximate\" and link the official page.", size=10.5, color=GREY)

    # privacy box
    t = doc.add_table(rows=1, cols=1)
    c = t.rows[0].cells[0]
    shade(c, 'FFF4D6')
    set_cell_borders_none(c)
    c.text = ""
    run(c.paragraphs[0], "Before you write: ", bold=True, size=10.5, color=NAVY)
    run(c.paragraphs[0], PRIVACY, size=10.5)
    doc.add_paragraph()

    heading(doc, "1. Post details", 1)
    rows = list(COMMON_FIELDS)
    if key == 'opportunities':
        rows = rows[:5] + OPPORTUNITY_FIELDS + rows[5:]
    fields_table(doc, rows)

    heading(doc, "2. Post text", 1)
    hint(doc, "Keep these headings in this order. Each heading becomes a section on the website.")
    for h, note in sections:
        heading(doc, h, 2)
        if note:
            hint(doc, note)
        body_placeholder(doc)

    heading(doc, "3. Before sending", 1)
    for line in (
        "Every figure is marked approximate and has a year.",
        "Every official claim has a link to the official page, not a news article or an agent's site.",
        "Nothing in the draft is specific to one school. The site speaks to students across India.",
        "A second counsellor has checked the facts and their name is in \"Checked by\".",
        "No private information anywhere in the draft.",
    ):
        p = doc.add_paragraph(style='List Bullet')
        run(p, line, size=10.5)

    os.makedirs(OUT, exist_ok=True)
    path = os.path.join(OUT, f"FP-post-template-{key}.docx")
    doc.save(path)
    print("wrote", os.path.relpath(path))


if __name__ == '__main__':
    for key, (label, sections) in BODIES.items():
        make(key, label, sections)
